#!/usr/bin/env python3
"""Generate bangkee-srs.docx — a comprehensive Software Requirements Specification.

Built with plain python-docx primitives only (headings, paragraphs, runs,
built-in table styles, page breaks). No hand-rolled OOXML / field codes /
custom cell shading — those open in Word but are rejected by Google Docs
("File could not open"). This mirrors the proven approach in gen_srs2.py.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

PRIMARY = RGBColor(0x00, 0x3D, 0x9B)
ACCENT = RGBColor(0x00, 0x6C, 0x47)

# Light heading colours via the built-in styles (safe; no XML surgery).
for lvl, color in [(1, PRIMARY), (2, PRIMARY), (3, ACCENT)]:
    try:
        doc.styles[f"Heading {lvl}"].font.color.rgb = color
    except Exception:
        pass


# ---------------- helpers (same shape as the reference script) ----------------
def h(text, level=1):
    return doc.add_heading(text, level=level)


def para(text="", bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def lead(label, text):
    p = doc.add_paragraph()
    p.add_run(label + " ").bold = True
    p.add_run(text)
    return p


def bullet(text, sub=False):
    p = doc.add_paragraph(style="List Bullet 2" if sub else "List Bullet")
    p.add_run(text)
    return p


def kv_bullet(label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(label + " — ").bold = True
    p.add_run(text)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hc = t.rows[0].cells
    for i, hh in enumerate(headers):
        hc[i].text = ""
        run = hc[i].paragraphs[0].add_run(str(hh))
        run.bold = True
    for row in rows:
        rc = t.add_row().cells
        for i, val in enumerate(row):
            rc[i].text = str(val)
    doc.add_paragraph()
    return t


def fr(fr_id, name, desc, inputs=None, processing=None, outputs=None,
       validation=None, errors=None, priority="High"):
    p = doc.add_paragraph()
    r = p.add_run(f"{fr_id}  {name}")
    r.bold = True
    r.font.color.rgb = ACCENT
    rows = [["Requirement", desc], ["Priority", priority]]
    if inputs:
        rows.append(["Inputs", inputs])
    if processing:
        rows.append(["Processing", processing])
    if outputs:
        rows.append(["Outputs", outputs])
    if validation:
        rows.append(["Validation", validation])
    if errors:
        rows.append(["Error handling", errors])
    table(["Field", "Specification"], rows)


def use_case(uc_id, title, actor, stakeholders, pre, trigger, main,
             alts=None, exceptions=None, post=None, freq=None):
    h(f"{uc_id}  {title}", 3)
    rows = [["Primary actor", actor],
            ["Stakeholders & interests", stakeholders],
            ["Preconditions", pre],
            ["Trigger", trigger]]
    if freq:
        rows.append(["Frequency of use", freq])
    table(["Field", "Value"], rows)
    para("Main success scenario", bold=True)
    for s in main:
        numbered(s)
    if alts:
        para("Alternate flows", bold=True)
        for a in alts:
            bullet(a)
    if exceptions:
        para("Exception flows", bold=True)
        for e in exceptions:
            bullet(e)
    if post:
        para("Postconditions", bold=True)
        for pc in post:
            bullet(pc)


# =====================================================================
# TITLE PAGE
# =====================================================================
doc.add_heading("Software Requirements Specification", level=0)
para("Bangkee — A digital credit book for small shops in Bhutan", bold=True, align="center")
para('"Replace the handwritten shop notebook with a shared ledger both the shop owner '
     'and the customer can see in real time."', italic=True, align="center")
para("Version 2.1 · Status: Baselined (MVP) + Planned modules · Date: 22 June 2026", align="center")
para("Audience: Product · Engineering (backend, web, mobile) · QA · Design · Maintainers",
     align="center")
doc.add_paragraph()
para("This document specifies the functional and non-functional requirements of Bangkee: the "
     "Rails 8 backend and web application (bangkee-web) including its JSON API, and the Flutter "
     "mobile client (bangkee-app). It is the authoritative reference for engineering, QA, security "
     "review, and onboarding. Sections 1–15 and the appendices document the core ledger MVP "
     "completed on 22 June 2026 (the as-built 2.0 baseline). Sections 16 and 17 specify two "
     "Planned modules — customer bank-transfer settlement with proof, and platform subscription "
     "billing with plan enforcement — that are committed for the next release but are not yet "
     "implemented; each is clearly marked Planned.")
doc.add_page_break()

# =====================================================================
# DOCUMENT CONTROL
# =====================================================================
h("Document Control", 1)
h("Revision History", 2)
table(["Version", "Date", "Author", "Summary of changes"],
      [["0.1", "22 Jun 2026", "Engineering", "Initial draft skeleton."],
       ["1.0", "22 Jun 2026", "Engineering", "First baseline covering the core ledger MVP."],
       ["2.0", "22 Jun 2026", "Engineering",
        "Comprehensive expansion: full data dictionary, use cases, per-endpoint API spec, "
        "screen specifications, validation and business-rule catalogs, authorization matrix, "
        "and acceptance criteria."],
       ["2.1", "22 Jun 2026", "Engineering",
        "Added two Planned modules: §16 Customer Payments via Bank Transfer (proof upload + owner "
        "confirmation) and §17 Subscription, Billing & Plan Enforcement (Nu 200/month, manual "
        "approval, warning + 2-week grace then feature lockout). Added the Platform Administrator "
        "user class and supporting data model, rules, use cases, API, and audit events."]])
h("Approvals", 2)
table(["Role", "Name", "Signature", "Date"],
      [["Product Owner", "", "", ""], ["Engineering Lead", "", "", ""],
       ["QA Lead", "", "", ""], ["Stakeholder (shop representative)", "", "", ""]])
h("Distribution", 2)
bullet("Product management")
bullet("Engineering (backend, web, mobile)")
bullet("Quality assurance")
bullet("Design")
bullet("Future maintainers and onboarding engineers")
doc.add_page_break()

# =====================================================================
# TABLE OF CONTENTS
# =====================================================================
h("Table of Contents", 1)
toc = [
    "Introduction", "Overall Description", "System Architecture",
    "Data Model & Data Dictionary", "Domain Rules & Business Logic",
    "Functional Requirements", "Use Cases", "External Interface Requirements",
    "Non-Functional Requirements", "Mobile Application Requirements (Flutter)",
    "Validation & Error-Handling Catalogue", "Security Requirements",
    "Data Management, Migration & Seed Data", "Testing & Acceptance Criteria",
    "Constraints, Assumptions & Risks",
    "Customer Payments via Bank Transfer (Proof & Confirmation) — Planned",
    "Subscription, Billing & Plan Enforcement — Planned",
    "Future Enhancements (Out of Scope)",
    "Appendix A — Requirements Traceability Matrix",
    "Appendix B — Authorization Matrix", "Appendix C — Glossary",
    "Appendix D — Audit Event Catalogue", "Appendix E — State-Transition Tables",
    "Appendix F — Open Issues",
]
for i, t in enumerate(toc, 1):
    para(f"{i}. {t}")
doc.add_page_break()

# =====================================================================
# 1. INTRODUCTION
# =====================================================================
h("1. Introduction", 1)
h("1.1 Purpose", 2)
para("This Software Requirements Specification (SRS) defines the functional and non-functional "
     "requirements for Bangkee, a digital credit-book (“khata”) system for small shops in "
     "Bhutan. It is the authoritative description of what the system does, the rules it enforces, "
     "and the constraints under which it operates. It serves the development team, quality "
     "assurance, product stakeholders, designers, and future maintainers. This baseline "
     "(version 2.0) documents the core ledger MVP and the interfaces that support its two clients: "
     "a Rails 8 web application and a Flutter mobile application.")
h("1.2 Scope", 2)
para("Bangkee replaces the handwritten shop notebook with a shared digital ledger that both the "
     "shop owner and the customer can view in real time. A shop owner records credit (goods given "
     "on tab) and payments against each customer's account; the running balance, transaction "
     "history, and overdue status are always visible to both parties. The product is a monorepo "
     "with two clients sharing one backend:")
kv_bullet("bangkee-web", "a Rails 8 (Hotwire) web application providing the shop-owner and customer "
          "web experience and a token-authenticated JSON API under /api/v1.")
kv_bullet("bangkee-app", "a Flutter mobile application, role-aware for both shop owners and "
          "customers, that consumes the JSON API.")
kv_bullet("design", "a shared Material 3 “Trust Blue” design system that both clients follow.")
para("Throughout this document a positive balance means the customer owes the shop. The MVP "
     "deliberately excludes the capabilities enumerated in Section 18. Sections 16 and 17 specify "
     "two Planned post-MVP modules (bank-transfer settlement and subscription billing).")
h("1.3 Intended Audience and Reading Suggestions", 2)
table(["Reader", "Suggested sections"],
      [["Product / stakeholders", "1, 2, 7 (use cases), 16–17 (planned modules), 18 (future work)"],
       ["Backend engineers", "3, 4, 5, 6, 8.2 (API), 12 (security), 14"],
       ["Web/mobile engineers", "3, 6, 8.1 (screens), 10 (mobile), 11 (validation)"],
       ["QA engineers", "6, 7, 11, 14 (acceptance criteria), Appendix A"],
       ["Maintainers / onboarding", "All; begin with 2 and 3"]])
h("1.4 Definitions, Acronyms & Abbreviations", 2)
table(["Term", "Definition"],
      [["Credit book / Khata", "A running tally of goods bought on credit and repayments, traditionally kept in a paper notebook."],
       ["Shop", "A single retail business owned by one shop-owner user. In this version an owner has exactly one shop."],
       ["Account", "A per-customer ledger held by a shop. May or may not be linked to a registered customer user."],
       ["Customer (account)", "The person who owes/repays a shop, represented by an Account; may later link to a User."],
       ["Transaction", "A single ledger entry — either a credit or a payment."],
       ["Credit", "A transaction that increases the customer's balance owed to the shop."],
       ["Payment", "A transaction that decreases the customer's balance."],
       ["Itemized credit", "A credit whose amount is derived from a list of line items."],
       ["Line item", "An individual good on an itemized credit (name, quantity, unit price)."],
       ["Balance", "Cached net amount owed, in chetrum. Positive = owes; negative = in credit; zero = settled."],
       ["Running balance", "The balance after each transaction is applied, shown in the ledger."],
       ["Void", "Reversible deactivation of a transaction; excluded from balance and ledger, retained for audit."],
       ["Audit event", "An append-only record of a significant action (created, edited, voided, account_created, customer_joined)."],
       ["Overdue", "An account with a positive balance whose last activity predates the shop's credit-due window."],
       ["Credit-due window", "The number of days (credit_due_days, default 30) after which an inactive owed balance is overdue."],
       ["Invitation", "A tokenised link (/join/:token) by which a customer links to a shop-created account."],
       ["Ngultrum (Nu.)", "Bhutanese currency unit. 1 Ngultrum = 100 chetrum."],
       ["Chetrum", "Minor currency unit; all money is stored as integer chetrum in *_cents columns."],
       ["Session", "An authenticated login instance, by signed cookie (web) or bearer token (API)."],
       ["Bearer token", "An opaque session token presented in the API Authorization header."],
       ["Hotwire", "Rails' HTML-over-the-wire front-end stack (Turbo + Stimulus)."],
       ["SRS / API / MVP", "Software Requirements Specification / Application Programming Interface / Minimum Viable Product."],
       ["FR / NFR / BR / UC", "Functional Requirement / Non-Functional Requirement / Business Rule / Use Case (identifier prefixes)."]])
h("1.5 Document Conventions", 2)
bullet("The verb “shall” denotes a mandatory requirement; “should” denotes a recommendation.")
bullet("Functional requirements are identified FR-<module>.<n>; business rules BR-<n>; non-functional NFR-<n>; use cases UC-<n>.")
bullet("“Owner” means an authenticated shop-owner acting on their own shop; “Customer” means an authenticated customer.")
bullet("Monetary values are expressed in Ngultrum for the reader but stored and transmitted as integer chetrum.")
h("1.6 References", 2)
bullet("Bangkee monorepo README.md and bangkee-app/README.md.")
bullet("design/bangkee_digital_ledger/DESIGN.md — Material 3 “Trust Blue” design system and screen mockups.")
bullet("Ruby on Rails 8.1 Guides — Active Record, built-in Authentication, Hotwire, Action Mailer.")
bullet("Flutter / Dart documentation; the provider and http packages.")
bullet("IEEE 830-1998 and ISO/IEC/IEEE 29148 — requirements specification structure references.")
h("1.7 Overview", 2)
para("Section 2 describes the product context, users, and constraints. Section 3 presents the "
     "system architecture. Section 4 is the complete data dictionary. Section 5 catalogues domain "
     "rules. Section 6 enumerates functional requirements. Section 7 provides detailed use cases. "
     "Section 8 covers external interfaces including a per-endpoint API specification and screen "
     "specifications. Section 9 lists non-functional requirements. Sections 10–14 cover the "
     "mobile app, validation, security, data/seed handling, and testing/acceptance. Sections 16 "
     "and 17 specify two Planned modules (bank-transfer settlement; subscription billing). "
     "Section 18 lists deferred work, and the appendices provide traceability, an authorization "
     "matrix, a glossary, the audit catalogue, and state-transition tables.")
doc.add_page_break()

# =====================================================================
# 2. OVERALL DESCRIPTION
# =====================================================================
h("2. Overall Description", 1)
h("2.1 Product Perspective", 2)
para("Bangkee is a new, self-contained client–server system. A single Rails 8 backend owns the "
     "PostgreSQL data store and exposes two surfaces: server-rendered HTML (Hotwire) for browser "
     "users, and a versioned JSON API under /api/v1 for the Flutter mobile clients. There is no "
     "third-party ledger or accounting system to integrate with in the MVP. The two clients share "
     "one data model, one set of domain rules, and — on the server — one set of serializers.")
h("2.2 Product Functions (Summary)", 2)
for f in [
    "Register and authenticate users as either shop owners or customers.",
    "Provision a shop automatically when a shop owner signs up.",
    "Create, search, filter, edit, and remove per-customer accounts within a shop.",
    "Record simple and itemized credits, and record payments with optional method and notes.",
    "Maintain a cached, always-current balance and a chronological ledger with running balances per account.",
    "Void and edit transactions while preserving an immutable audit trail.",
    "Surface role-aware dashboards (shop totals and overdue list; customer's debts across shops).",
    "Generate per-account statements.",
    "Invite customers via a tokenised link and link them to their existing account.",
    "Track overdue accounts against a configurable credit-due window.",
    "Serve all of the above to the mobile app through a token-authenticated JSON API.",
]:
    bullet(f)
h("2.3 User Classes and Characteristics", 2)
table(["User class", "Role", "Characteristics & privileges"],
      [["Shop owner", "role = 0 (shop_owner)",
        "Owns exactly one shop. Full read/write over their shop's accounts and transactions: add/edit/remove customers, record credits and payments, edit and void transactions, view the shop dashboard, overdue list, and statements. Cannot access any other shop's data."],
       ["Customer", "role = 1 (customer; default)",
        "May hold accounts across multiple shops. Read-only: views their own ledgers, balances, statements, and the customer dashboard. Cannot create, edit, or void transactions, and cannot view shop-wide aggregates."],
       ["Unauthenticated visitor", "—",
        "May reach the login, registration, password-reset, and invitation (/join/:token) pages only. All other pages redirect to login."],
       ["Platform administrator (Planned, §17)", "platform_admin flag",
        "The Bangkee app owner/operator. Reviews and approves/rejects shop subscription payments, views all shops and their subscription status, and may enable/disable a shop's plan. Has no access to any shop's customer ledger data; operates only on billing/subscription records."]])
h("2.4 Operating Environment", 2)
table(["Component", "Technology / version"],
      [["Backend framework", "Ruby on Rails ~> 8.1.3 on Ruby 4.0.0"],
       ["Database", "PostgreSQL (uses jsonb metadata, compound and unique indexes)"],
       ["Web front-end", "Hotwire (Turbo + Stimulus); Propshaft asset pipeline; importmap-rails"],
       ["Browser policy", "Modern browsers only (allow_browser versions: :modern)"],
       ["PWA", "Web manifest and service worker present (app/views/pwa)"],
       ["Mobile client", "Flutter (Dart SDK >= 3.4.0 < 4.0.0); http, provider, shared_preferences, intl"],
       ["Auth (web)", "Rails 8 built-in auth: sessions + has_secure_password (bcrypt); signed httponly cookies"],
       ["Auth (API)", "Opaque per-session bearer tokens"],
       ["Email", "Action Mailer (password reset), delivered asynchronously (deliver_later)"],
       ["Design system", "Material 3 “Trust Blue”"],
       ["Deployment", "Containerised (Dockerfile present)"],
       ["Health check", "GET /up (Rails health endpoint)"]])
h("2.5 Design and Implementation Constraints", 2)
kv_bullet("Money as integers", "All monetary amounts are stored as integer chetrum in *_cents "
          "columns. Major-unit (Ngultrum) values are derived only at the form/display layer via the "
          "HasMoneyAttribute concern. No floating-point money is persisted.")
kv_bullet("Built-in authentication", "Authentication uses Rails 8's built-in mechanism (not Devise). "
          "A single User model carries a role enum; sign-up provisions a shop for owners.")
kv_bullet("Association naming", "The LineItem→Transaction association is named :purchase "
          "(class_name 'Transaction', foreign key transaction_id) to avoid colliding with Active "
          "Record's built-in transaction method. This naming must not be “corrected.”")
kv_bullet("Void, never hard-delete ledger entries", "Transactions are voided (soft-deactivated) "
          "rather than deleted in normal operation, so the audit trail and history remain intact.")
kv_bullet("Single shop per owner", "This version supports exactly one shop per owner; multi-shop is deferred.")
kv_bullet("Cached, recomputed balance", "Each account caches balance_cents but always recomputes it "
          "from source transactions, never by incremental mutation, to prevent drift.")
kv_bullet("Design fidelity", "Both clients must adhere to the shared Material 3 “Trust Blue” tokens and screen designs.")
h("2.6 User Documentation", 2)
bullet("Monorepo README.md — setup, demo logins, and an API endpoint summary.")
bullet("bangkee-app/README.md — mobile build/run instructions and project structure.")
bullet("design/bangkee_digital_ledger/DESIGN.md — design tokens and screen mockups.")
bullet("In-product guidance via flash notices and form-validation messages.")
h("2.7 Assumptions and Dependencies", 2)
bullet("Users have intermittent but generally available internet connectivity; offline operation is out of scope for the MVP.")
bullet("Each shop owner manages a single shop with a modest number of customer accounts.")
bullet("Phone numbers are optional and are not verified in this version.")
bullet("The default credit-due window is 30 days and is configurable per shop (must be > 0).")
bullet("The currency is the Bhutanese Ngultrum; localisation to Dzongkha is deferred.")
bullet("Email delivery infrastructure is available for password-reset messages.")
h("2.8 Apportioning of Requirements", 2)
para("Requirements describe the current baseline unless explicitly marked as Planned (§16–§17) or "
     "deferred (§18). Sections 1–15 reflect the as-built 2.0 system. Section 18 "
     "lists capabilities intentionally postponed beyond the MVP; those are out of scope for "
     "verification against this baseline.")
doc.add_page_break()

# =====================================================================
# 3. ARCHITECTURE
# =====================================================================
h("3. System Architecture", 1)
h("3.1 Architectural Overview", 2)
para("Bangkee follows a conventional Rails Model-View-Controller architecture, extended with a "
     "parallel API controller namespace. The same models and domain logic back both surfaces:")
kv_bullet("Models", "User, Shop, Account, Transaction, LineItem, AuditEvent, Session, plus the "
          "Current request-scoped attributes object. Cross-cutting behaviour lives in concerns "
          "(HasMoneyAttribute, Auditable, Authentication).")
kv_bullet("Web controllers", "Server-render Hotwire HTML and enforce authentication/authorization "
          "via the Authentication concern and ApplicationController helpers.")
kv_bullet("API controllers", "Api::V1::* inherit from a BaseController (ActionController::API) that "
          "does bearer-token authentication and shared error handling; responses use ApiSerializers.")
kv_bullet("Views", "ERB templates plus partials (account form, credit/line-item fields, transaction row).")
h("3.2 Request Lifecycle", 2)
para("Web requests resume a session from a signed cookie; unauthenticated requests to protected "
     "actions are redirected to login with a return-to URL. API requests authenticate by matching "
     "the bearer token to a Session and loading its user; failures return 401. The request-scoped "
     "Current.session (and delegated Current.user) make the actor available to models for auditing.")
h("3.3 Data Flow for a Transaction", 2)
numbered("An owner submits a credit or payment (web form or API request).")
numbered("The controller builds a Transaction on the account, sets its kind and created_by, and saves.")
numbered("On create, the model defaults occurred_at and, for itemized credits, derives amount_cents from the line items.")
numbered("Validations ensure amount_cents > 0 and occurred_at is present.")
numbered("An after_save callback calls Account#recompute_balance!, recomputing balance_cents and last_activity_at from active transactions.")
numbered("The controller writes a “created” audit event and renders the updated account (HTML redirect or JSON).")
h("3.4 Deployment View", 2)
bullet("A containerised Rails application server fronting a PostgreSQL database.")
bullet("Browser clients connect over HTTP/HTTPS and receive server-rendered HTML with Hotwire.")
bullet("Mobile clients connect to /api/v1 over HTTP/HTTPS using a compile-time configurable base URL (default http://10.0.2.2:3007/api/v1 for the Android emulator).")
bullet("A health endpoint (/up) supports external uptime monitoring.")
doc.add_page_break()

# =====================================================================
# 4. DATA DICTIONARY
# =====================================================================
h("4. Data Model & Data Dictionary", 1)
para("All tables use a bigint primary key id and created_at / updated_at timestamps unless noted. "
     "Monetary columns are bigint chetrum. The following sub-sections document every persisted "
     "column, its constraints, and its meaning.")
h("4.1 Entity-Relationship Summary", 2)
bullet("A User (shop_owner) owns one Shop (owned_shops, dependent: destroy).")
bullet("A Shop has many Accounts; a Shop has many Transactions through Accounts.")
bullet("An Account belongs to a Shop and optionally to a customer User; it has many Transactions.")
bullet("A Transaction belongs to an Account, a creator User (created_by), and optionally a voider User (voided_by); it has many LineItems (named :purchase).")
bullet("A LineItem belongs to its Transaction via the :purchase association (foreign key transaction_id).")
bullet("AuditEvent is polymorphic (auditable) and optionally references an actor User.")
bullet("A Session belongs to a User.")
h("4.2 users", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["email_address", "string", "no", "—", "Unique; normalized (trimmed, lower-cased)"],
       ["password_digest", "string", "no", "—", "bcrypt digest (has_secure_password)"],
       ["name", "string", "no", "—", "Display name; presence validated"],
       ["phone", "string", "yes", "—", "Normalized (trimmed, blank→null); indexed"],
       ["role", "integer", "no", "1", "Enum: 0 shop_owner, 1 customer; validated"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("Indexes: unique on email_address; index on phone.")
h("4.3 shops", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["name", "string", "no", "—", "Shop name; presence validated"],
       ["owner_id", "bigint", "no", "—", "FK → users; the shop_owner"],
       ["credit_due_days", "integer", "no", "30", "Overdue window in days; must be > 0"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("Indexes: index on owner_id. Foreign key: shops.owner_id → users.id.")
h("4.4 accounts", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["shop_id", "bigint", "no", "—", "FK → shops"],
       ["customer_id", "bigint", "yes", "—", "FK → users; null until the customer joins"],
       ["customer_name", "string", "no", "—", "Display name set by owner; presence validated"],
       ["customer_phone", "string", "yes", "—", "Normalized (blank→null)"],
       ["balance_cents", "bigint", "no", "0", "Cached net owed (chetrum); positive = owes"],
       ["invite_token", "string", "no", "—", "Unique; auto-generated (has_secure_token)"],
       ["last_activity_at", "datetime", "yes", "—", "Max occurred_at of active txns; drives overdue"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("Indexes: unique on invite_token; index on shop_id and customer_id; compound on "
     "(shop_id, customer_id) and (shop_id, customer_phone). FKs to shops and users (customer).")
h("4.5 transactions", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["account_id", "bigint", "no", "—", "FK → accounts"],
       ["kind", "integer", "no", "0", "Enum: 0 credit, 1 payment; validated"],
       ["amount_cents", "bigint", "no", "0", "Amount in chetrum; must be > 0"],
       ["description", "string", "yes", "—", "Free-text label (credits)"],
       ["notes", "text", "yes", "—", "Free-text note"],
       ["payment_method", "string", "yes", "—", "e.g. Cash, Mobile wallet (payments)"],
       ["itemized", "boolean", "no", "false", "True when amount derives from line items"],
       ["occurred_at", "datetime", "no", "—", "Business time; defaults to now on create"],
       ["created_by_id", "bigint", "no", "—", "FK → users; recorder"],
       ["voided_at", "datetime", "yes", "—", "Set when voided; excludes from balance/ledger"],
       ["voided_by_id", "bigint", "yes", "—", "FK → users; who voided"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("Indexes: index on account_id; compound on (account_id, occurred_at) and (account_id, kind); "
     "index on created_by_id and voided_by_id. FKs to accounts and users (created_by, voided_by).")
h("4.6 line_items", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["transaction_id", "bigint", "no", "—", "FK → transactions (the :purchase)"],
       ["name", "string", "no", "—", "Item name; presence validated"],
       ["quantity", "decimal(10,2)", "no", "1.0", "Must be > 0"],
       ["unit_price_cents", "bigint", "no", "0", "Per-unit price (chetrum); must be >= 0"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("total_cents = round(quantity × unit_price_cents). Index on transaction_id; FK to transactions.")
h("4.7 audit_events", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["auditable_type", "string", "no", "—", "Polymorphic type (e.g. Account, Transaction)"],
       ["auditable_id", "bigint", "no", "—", "Polymorphic id"],
       ["actor_id", "bigint", "yes", "—", "FK → users; null ⇒ “System”"],
       ["action", "string", "no", "—", "e.g. created, edited, voided, account_created, customer_joined"],
       ["metadata", "jsonb", "no", "{}", "Per-action structured detail"],
       ["created_at", "datetime", "no", "—", "Timestamp (append-only)"]])
para("Indexes: on (auditable_type, auditable_id) and (auditable_type, auditable_id, created_at); "
     "index on actor_id. FK: audit_events.actor_id → users.id.")
h("4.8 sessions", 2)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["user_id", "bigint", "no", "—", "FK → users"],
       ["token", "string", "yes", "—", "Unique; bearer token for API auth"],
       ["ip_address", "string", "yes", "—", "Captured at creation"],
       ["user_agent", "string", "yes", "—", "Captured at creation"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("Indexes: unique on token; index on user_id. FK: sessions.user_id → users.id.")
h("4.9 Enumerations", 2)
table(["Enum", "Model", "Values"],
      [["role", "User", "shop_owner = 0, customer = 1 (default 1); validated"],
       ["kind", "Transaction", "credit = 0 (default), payment = 1; validated"]])
doc.add_page_break()

# =====================================================================
# 5. DOMAIN RULES
# =====================================================================
h("5. Domain Rules & Business Logic", 1)
para("Business rules are identified BR-n and are enforced in the models and concerns so that they "
     "hold identically for the web and API surfaces.")
h("5.1 Money Representation", 2)
table(["ID", "Rule"],
      [["BR-1", "All monetary amounts are persisted as integer chetrum in *_cents columns (bigint)."],
       ["BR-2", "Major-unit accessors (amount, unit_price) convert to/from BigDecimal, multiplying by 100 and rounding on write; blank or non-numeric input maps to nil."],
       ["BR-3", "Amounts display in Ngultrum via the ngultrum helper: whole values drop decimals (“Nu. 1,250”); fractional values keep two (“Nu. 1,250.50”), with comma thousands separators."],
       ["BR-4", "The API transmits amounts as integer chetrum (*_cents); clients format to “Nu.” themselves."]])
h("5.2 Balance & Ledger", 2)
table(["ID", "Rule"],
      [["BR-5", "A positive balance means the customer owes the shop; negative means in credit; zero means settled."],
       ["BR-6", "balance_cents = Σ(active credits.amount_cents) − Σ(active payments.amount_cents)."],
       ["BR-7", "Account#recompute_balance! recomputes balance_cents and last_activity_at from active transactions in Transaction after_save and after_destroy."],
       ["BR-8", "last_activity_at is the maximum occurred_at over active transactions."],
       ["BR-9", "A credit contributes +amount_cents and a payment −amount_cents to the running balance (signed_amount_cents)."],
       ["BR-10", "Account#ledger returns active transactions in chronological order (occurred_at, then id), each annotated with the running balance after it is applied; views may show it most-recent-first."],
       ["BR-11", "Voided transactions are excluded from both balance and ledger."]])
h("5.3 Itemized Credits", 2)
table(["ID", "Rule"],
      [["BR-12", "When a credit is itemized with at least one non-deleted line item, amount_cents is set to the sum of line-item totals before validation (sync_itemized_total)."],
       ["BR-13", "A line item's total_cents = round(quantity × unit_price_cents); quantity must be > 0 and unit_price_cents >= 0."],
       ["BR-14", "Nested line-item attributes with both a blank name and blank unit price are rejected (not persisted)."]])
h("5.4 Void & Audit", 2)
table(["ID", "Rule"],
      [["BR-15", "Transaction#void!(by:) sets voided_at and voided_by inside a database transaction and is a no-op (returns false) if already voided."],
       ["BR-16", "Voiding writes a “voided” audit event capturing amount_cents and kind; balance recomputation then removes the entry's effect."],
       ["BR-17", "Audit events are append-only and written explicitly with per-action metadata (Auditable#log_audit), not as raw column diffs."],
       ["BR-18", "Auditable actions in the MVP: account_created, created (credit/payment), edited, voided, customer_joined. An event with no actor displays as “System.”"]])
h("5.5 Overdue Determination", 2)
table(["ID", "Rule"],
      [["BR-19", "An account is overdue when balance_cents > 0 and last_activity_at is null or older than credit_due_days days before now."],
       ["BR-20", "Shop#overdue_accounts returns owed accounts whose last_activity_at is null or before the threshold; credit_due_days defaults to 30 and must be > 0."],
       ["BR-21", "Settled or in-credit accounts (balance_cents <= 0) are never overdue."]])
h("5.6 Account Linking (Invitations)", 2)
table(["ID", "Rule"],
      [["BR-22", "Each account has a unique, auto-generated invite_token. joined? is true once customer_id is set."],
       ["BR-23", "Visiting /join/:token while signed out stashes the token and routes the visitor to customer registration."],
       ["BR-24", "Visiting while signed in as a customer links the account to that customer, backfilling customer_phone from the user if absent, and logs customer_joined."],
       ["BR-25", "An already-claimed account (by a different customer) or a shop-owner visitor is refused with an explanatory message."],
       ["BR-26", "On registration, a pending invite token (web session) or supplied invite_token (API) is claimed for new customers."]])
doc.add_page_break()

# =====================================================================
# 6. FUNCTIONAL REQUIREMENTS
# =====================================================================
h("6. Functional Requirements", 1)
para("Each requirement specifies its behaviour and, where useful, inputs, processing, outputs, "
     "validation, and error handling. Priorities: High = MVP-critical, Medium = important, Low = desirable.")
h("6.1 Authentication & Registration", 2)
fr("FR-1.1", "User registration",
   "The system shall allow a visitor to register as a shop owner or customer.",
   inputs="name, email_address, phone (optional), password, password_confirmation, role; optional shop_name (owners).",
   processing="Create the User; if shop_owner, create an owned shop named shop_name or “<name>'s Shop”; claim any pending invitation for customers; start a session.",
   outputs="Authenticated session and redirect (web) / { token, user } (API).",
   validation="email_address unique and normalized; name present; role in {shop_owner, customer}; password confirmation must match (web).",
   errors="On failure, re-render the form with messages (web, 422) or return { error, details } (API, 422).")
fr("FR-1.2", "Shop provisioning on owner sign-up",
   "On successful shop-owner registration, the system shall create exactly one shop owned by that user.")
fr("FR-1.3", "Login",
   "The system shall authenticate by email and password and establish a session.",
   inputs="email_address, password.",
   processing="User.authenticate_by; on success create a Session capturing user_agent and ip_address.",
   outputs="Signed cookie session + redirect to return-to URL (web); { token, user } (API).",
   errors="Invalid credentials redirect to login with an alert (web) / 401 { error } (API).")
fr("FR-1.4", "Login rate limiting",
   "The system shall rate-limit web login and password-reset to 10 attempts within 3 minutes, responding with a “try again later” message when exceeded.",
   priority="Medium")
fr("FR-1.5", "Logout",
   "The system shall allow users to log out, destroying the current session (cookie cleared on web; 204 on API).")
fr("FR-1.6", "Password reset",
   "The system shall provide a password-reset flow keyed by a signed, expiring token.",
   processing="Request emails reset instructions if the address exists (no account enumeration); update sets the new password and destroys all of the user's sessions.",
   validation="Token signature valid and unexpired; password confirmation matches.",
   errors="Invalid/expired token redirects to the reset request with an alert; mismatched passwords re-prompt.")
fr("FR-1.7", "Authentication gate & return-to",
   "Unauthenticated access to protected pages shall redirect to login and, after authentication, return the user to the originally requested URL.")
fr("FR-1.8", "Modern-browser policy",
   "The web application shall serve only modern browsers (webp, import maps, CSS :has/nesting).",
   priority="Low")
h("6.2 Shop & Customer Account Management", 2)
fr("FR-2.1", "Create customer account",
   "An owner shall create a customer account in their shop with a required customer_name and optional customer_phone; the system assigns a unique invite_token and audits account_created.",
   validation="customer_name present; customer_phone normalized (blank→null).")
fr("FR-2.2", "List, search, and filter accounts",
   "An owner shall view their accounts, searchable by name or phone (case-insensitive) and filterable by all / outstanding / overdue / settled, ordered by customer name.",
   processing="search matches customer_name/customer_phone via ILIKE; filters map to outstanding (balance>0), overdue (Shop#overdue_accounts), settled (balance<=0).")
fr("FR-2.3", "View account detail",
   "The system shall show an account's ledger (most-recent-first) and the 20 most recent audit events.",
   validation="Viewer must be the shop owner or the linked customer (else redirect/403).")
fr("FR-2.4", "Edit account", "An owner shall edit a customer's name and phone within their own shop.")
fr("FR-2.5", "Remove account",
   "An owner shall remove (destroy) a customer account in their own shop; dependent transactions, line items, and audit events are destroyed.",
   priority="Medium")
fr("FR-2.6", "Shop-scoped access",
   "The system shall scope all account reads and writes to the requesting owner's shop; cross-shop access is forbidden.")
h("6.3 Transactions — Credits", 2)
fr("FR-3.1", "Record a simple credit",
   "An owner shall record a credit as a single amount with optional description, notes, and occurred_at.",
   inputs="amount, description, notes, occurred_at.",
   validation="amount_cents > 0; occurred_at present (defaults to now).",
   outputs="Credit persisted; balance recomputed; “created” audit event; confirmation showing the formatted amount.")
fr("FR-3.2", "Record an itemized credit",
   "An owner shall record an itemized credit composed of line items; the total is derived from the line items.",
   inputs="itemized=true; line_items[] of { name, quantity, unit_price }.",
   processing="amount_cents = Σ round(quantity × unit_price_cents) over non-deleted line items.",
   validation="Each line item: name present, quantity > 0, unit_price_cents >= 0; blank name+price rows are rejected.")
fr("FR-3.3", "Edit a credit",
   "An owner shall edit a credit (amount or line items, description, notes, occurred_at); the edit is audited and the balance recomputed.",
   priority="Medium")
h("6.4 Transactions — Payments", 2)
fr("FR-4.1", "Record a payment",
   "An owner shall record a payment with an amount and optional payment_method, notes, and occurred_at.",
   validation="amount_cents > 0; occurred_at present.",
   outputs="Payment persisted; balance decreased; “created” audit event; confirmation showing the formatted amount.")
fr("FR-4.2", "Edit a payment",
   "An owner shall edit a payment's amount, method, notes, or occurred_at; audited and balance recomputed.",
   priority="Medium")
h("6.5 Transactions — Void & Recompute", 2)
fr("FR-5.1", "Void a transaction",
   "An owner shall void a transaction, marking it voided (voided_at, voided_by), excluding it from balance and ledger, and writing a “voided” audit event.",
   errors="Voiding an already-voided transaction is a no-op and reports “already voided” (alert / 422).")
fr("FR-5.2", "Automatic balance recomputation",
   "After any transaction create, update, void, or destroy, the system shall recompute the account's cached balance and last-activity timestamp from active transactions.")
fr("FR-5.3", "Transaction ownership",
   "Only the owner of the transaction's shop may edit or void it; other users are denied (redirect / 403).")
h("6.6 Dashboards", 2)
fr("FR-6.1", "Role-aware dashboard",
   "The root path shall present an owner or customer dashboard according to the user's role.")
fr("FR-6.2", "Owner dashboard",
   "The owner dashboard shall show total outstanding, customer count, the overdue accounts list, repayments received this month, and up to 8 recent active transactions.")
fr("FR-6.3", "Customer dashboard",
   "The customer dashboard shall show total owed across all linked shops, amount spent this month, the customer's accounts with their shops, and up to 8 recent active transactions.")
h("6.7 Statements & Ledger", 2)
fr("FR-7.1", "Per-account statement",
   "The system shall provide a per-account statement view summarising the chronological ledger with a generation timestamp.",
   validation="Viewer must be the shop owner or the linked customer.")
fr("FR-7.2", "Running balance presentation",
   "The ledger shall present active transactions with the running balance after each entry.")
h("6.8 Invitations & Linking", 2)
fr("FR-8.1", "Invitation link",
   "The system shall expose /join/:token reachable without prior authentication while still recognising an already-logged-in customer.")
fr("FR-8.2", "Deferred linking via registration",
   "If no user is signed in, the system shall stash the invite token and route the visitor to customer registration, linking the account on sign-up.")
fr("FR-8.3", "Immediate linking",
   "If a customer is signed in, the system shall link the account to them (backfilling phone if absent) and audit customer_joined.")
fr("FR-8.4", "Invitation guards",
   "If the account is already claimed by another customer, or a shop owner opens the link, the system shall refuse with an explanatory message.")
h("6.9 Overdue Tracking", 2)
fr("FR-9.1", "Overdue definition",
   "An account shall be overdue when it has a positive balance and its last activity is null or older than the shop's credit-due window.")
fr("FR-9.2", "Overdue surfacing",
   "Owners shall be able to filter accounts by overdue and see overdue accounts on the dashboard.")
fr("FR-9.3", "Configurable window",
   "Each shop shall have a configurable credit_due_days (default 30, must be > 0).")
h("6.10 JSON API", 2)
para("See Section 8.2 for the full per-endpoint specification. The following are cross-cutting API requirements.")
fr("FR-10.1", "Bearer authentication",
   "The API shall require a valid bearer token for all endpoints except session create and registration create, returning 401 otherwise.")
fr("FR-10.2", "Owner authorization",
   "Owner-only API endpoints shall return 403 for non-owners; account access shall be limited to the shop owner or the linked customer.")
fr("FR-10.3", "Consistent errors",
   "The API shall return 404 for missing records and 422 with an error message and details for validation failures.")
fr("FR-10.4", "Shared serialization",
   "API responses shall use shared serializers so payloads stay consistent with the domain model (amounts as integer chetrum).")
doc.add_page_break()

# =====================================================================
# 7. USE CASES
# =====================================================================
h("7. Use Cases", 1)
para("Use cases describe goal-oriented interactions. Each lists the primary actor, stakeholders, "
     "preconditions, trigger, the main success scenario, and notable alternates/exceptions.")
use_case("UC-1", "Shop owner signs up and opens a shop",
         "Visitor (becomes shop owner)",
         "Owner wants a working shop ledger; the business wants accurate provisioning.",
         "Not authenticated.", "Visitor chooses “Sign up” as a shop owner.",
         ["Visitor submits name, email, password, and optional shop name.",
          "System validates and creates the user with role shop_owner.",
          "System creates a shop owned by the user.",
          "System starts a session and redirects to the owner dashboard."],
         alts=["1a. Visitor omits a shop name → the shop is named “<name>'s Shop”."],
         exceptions=["2a. Email already in use or invalid → form re-rendered with errors (422)."],
         post=["A shop owner and their shop exist; the owner is logged in."], freq="Once per owner")
use_case("UC-2", "Owner adds a customer account", "Shop owner",
         "Owner needs to track a customer's credit; customer will later join.",
         "Authenticated as a shop owner.", "Owner chooses “Add customer.”",
         ["Owner enters customer name and optional phone.",
          "System creates the account with a unique invite token.",
          "System records an account_created audit event.",
          "System shows the new account detail page."],
         exceptions=["1a. Name blank → form re-rendered with a validation error."],
         post=["A customer account exists, ready for transactions and invitation."], freq="Frequent")
use_case("UC-3", "Owner records a simple credit", "Shop owner",
         "Owner gives goods on tab and needs the balance updated.",
         "Authenticated owner; the account exists in the owner's shop.", "Owner chooses “Record credit.”",
         ["Owner enters an amount and optional description/notes.",
          "System validates amount > 0 and defaults occurred_at to now.",
          "System saves the credit and recomputes the account balance.",
          "System logs a “created” audit event and confirms with the formatted amount."],
         exceptions=["2a. Amount is zero or non-numeric → form re-rendered with an error (422)."],
         post=["Balance increased by the credit amount; ledger updated."], freq="Very frequent")
use_case("UC-4", "Owner records an itemized credit", "Shop owner",
         "Owner sells several items and wants an itemized record.",
         "Authenticated owner; account exists.", "Owner chooses “Record credit” in itemized mode.",
         ["Owner adds line items (name, quantity, unit price).",
          "System derives the total from the line items.",
          "System validates each line item and the derived total > 0.",
          "System saves the credit and line items, recomputes the balance, and confirms."],
         alts=["1a. Owner adds or removes rows dynamically before saving.",
               "1b. Rows with blank name and price are ignored."],
         exceptions=["3a. A line item has quantity <= 0 → validation error."],
         post=["Balance increased by the sum of line-item totals."], freq="Frequent")
use_case("UC-5", "Owner records a payment", "Shop owner",
         "Customer repays; owner needs the balance reduced.",
         "Authenticated owner; account exists.", "Owner chooses “Record payment.”",
         ["Owner enters an amount and optional method/notes.",
          "System validates amount > 0 and saves the payment.",
          "System recomputes the balance (decreased) and logs the event."],
         exceptions=["2a. Amount invalid → form re-rendered with an error."],
         post=["Balance decreased by the payment amount."], freq="Very frequent")
use_case("UC-6", "Owner voids a mistaken transaction", "Shop owner",
         "Owner needs to correct an error without losing history.",
         "Authenticated owner; a non-voided transaction exists in their shop.", "Owner chooses “Void” on a transaction.",
         ["System marks the transaction voided with voider and time.",
          "System writes a “voided” audit event capturing amount and kind.",
          "System recomputes the balance, excluding the voided entry.",
          "System confirms the void."],
         exceptions=["1a. Transaction already voided → no-op with an “already voided” message."],
         post=["The transaction is excluded from balance and ledger but retained for audit."], freq="Occasional")
use_case("UC-7", "Owner reviews overdue customers", "Shop owner",
         "Owner wants to follow up on stale debts.", "Authenticated owner.",
         "Owner opens the dashboard or filters accounts by “overdue.”",
         ["System computes overdue accounts (owed balance, stale last activity).",
          "System lists them ordered by customer name.",
          "Owner opens an account to review its ledger and follow up."],
         post=["Owner has the current overdue list."], freq="Regular")
use_case("UC-8", "Customer joins a shop via invitation", "Customer",
         "Customer wants visibility of what they owe.",
         "Owner created an account and shared /join/:token.", "Customer opens the invitation link.",
         ["If signed out, system stashes the token and routes to customer sign-up.",
          "Customer registers (or is already signed in as a customer).",
          "System links the account to the customer, backfilling phone if absent.",
          "System logs customer_joined and shows the account."],
         alts=["2a. Already signed in as a customer → immediate linking."],
         exceptions=["3a. Account already claimed by another customer → refused.",
                     "3b. A shop owner opens the link → refused with guidance."],
         post=["The customer is linked to the account and can view it."], freq="Once per customer per shop")
use_case("UC-9", "Customer reviews balances across shops", "Customer",
         "Customer wants to know total owed and recent activity.",
         "Authenticated as a customer with at least one linked account.", "Customer opens their dashboard.",
         ["System shows total owed across shops and spend this month.",
          "System lists linked accounts with shops and recent transactions.",
          "Customer opens an account to view its full ledger or statement."],
         post=["Customer has an up-to-date view of their debts."], freq="Regular")
use_case("UC-10", "Generate a statement", "Shop owner or linked customer",
         "Both parties want a clear chronological statement.",
         "Authenticated and authorized for the account.", "User opens the account statement.",
         ["System builds the chronological ledger with running balances.",
          "System renders the statement with a generation timestamp."],
         post=["A statement view is available for review."], freq="Occasional")
use_case("UC-11", "User resets a forgotten password", "Any registered user",
         "User must regain access securely.", "User knows their email.", "User requests a password reset.",
         ["System sends reset instructions if the email exists (no enumeration).",
          "User opens the emailed link and sets a new password.",
          "System updates the password and destroys all existing sessions.",
          "User logs in with the new password."],
         exceptions=["2a. Token invalid/expired → user is asked to request a new link.",
                     "3a. Password confirmation mismatch → re-prompt."],
         post=["Password changed; prior sessions invalidated."], freq="Rare")
use_case("UC-12", "Mobile owner records a credit via the API", "Shop owner (mobile app)",
         "Owner records credit on the go.", "App holds a valid bearer token for a shop owner.",
         "Owner submits a credit in the app.",
         ["App POSTs to /accounts/:id/credits with the transaction payload.",
          "Server authenticates the token and authorizes the owner.",
          "Server records the credit, recomputes the balance, and audits it.",
          "Server returns the created transaction as JSON (201)."],
         exceptions=["2a. Missing/invalid token → 401.", "2b. Non-owner → 403.",
                     "3a. Validation error → 422 with details."],
         post=["Credit recorded; app refreshes the account."], freq="Frequent")
doc.add_page_break()

# =====================================================================
# 8. EXTERNAL INTERFACES
# =====================================================================
h("8. External Interface Requirements", 1)
h("8.1 User Interface & Screen Specifications", 2)
para("The UI follows the Material 3 “Trust Blue” design system, anchored on a primary blue "
     "(#003D9B) with a green secondary (#006C47) and a red error (#BA1A1A) on light surfaces. "
     "Screen designs exist as mockups under design/. The table specifies the principal screens.")
table(["Screen", "Audience", "Primary content & actions"],
      [["Login", "All", "Email/password form; links to register and reset; rate-limited."],
       ["Register", "All", "Role-aware sign-up (owner adds shop name); claims pending invite."],
       ["Owner dashboard", "Owner", "Total outstanding, customer count, repayments this month, overdue list, recent transactions."],
       ["Customer list & search", "Owner", "Searchable, filterable account list (all/outstanding/overdue/settled); “Add customer.”"],
       ["Add/Edit customer", "Owner", "Name and phone form."],
       ["Account detail / ledger", "Owner & customer", "Balance, ledger with running balances, recent audit events; record credit/payment and void (owner)."],
       ["Record transaction", "Owner", "Simple or itemized credit; payment with method/notes."],
       ["Statement", "Owner & customer", "Chronological ledger with running balances and generation time."],
       ["Customer dashboard", "Customer", "Total owed across shops, spend this month, linked accounts, recent transactions."],
       ["Overdue management", "Owner", "Overdue accounts for follow-up."],
       ["Settings (mobile)", "All", "Profile and sign out."]])
lead("UI requirement.", "All screens shall use the Trust Blue tokens; balances shall be styled by "
     "state (owed / settled / in-credit) and amounts formatted in Ngultrum. Validation errors shall "
     "be shown inline on forms.")
h("8.2 JSON API Specification", 2)
para("Base path /api/v1. Authentication: Authorization: Bearer <token> (obtain via POST /session). "
     "Content type application/json. Amounts are integer chetrum. Standard error envelope: "
     "{ error: string, details?: string[] }.")
h("8.2.1 Endpoint summary", 3)
table(["Method & path", "Purpose", "Auth", "Success"],
      [["POST /session", "Log in", "Any", "201 { token, user }"],
       ["DELETE /session", "Log out", "Auth", "204"],
       ["POST /registration", "Sign up", "Any", "201 { token, user }"],
       ["GET /dashboard", "Role-aware dashboard", "Auth", "200 (role payload)"],
       ["GET /accounts?q=&filter=", "List customers", "Owner", "200 { accounts:[] }"],
       ["GET /accounts/:id", "Account detail + ledger", "Owner/linked customer", "200 (detail)"],
       ["POST /accounts", "Add customer", "Owner", "201 (detail)"],
       ["POST /accounts/:id/credits", "Record credit", "Owner", "201 (transaction)"],
       ["POST /accounts/:id/payments", "Record payment", "Owner", "201 (transaction)"],
       ["POST /transactions/:id/void", "Void a transaction", "Owner", "200 (transaction)"]])
h("8.2.2 Authentication & registration", 3)
lead("POST /session.", "Request: { email_address, password }. Success 201: { token, user }. "
     "Failure 401: { error: “Invalid email or password.” }.")
lead("DELETE /session.", "Destroys the current session; returns 204 No Content.")
lead("POST /registration.", "Request: { user: { name, email_address, phone, password, role }, "
     "shop_name?, invite_token? }. Creates the user; for owners creates a shop; for customers with "
     "invite_token links the account. Success 201: { token, user }. Failure 422: { error, details }.")
h("8.2.3 Dashboard", 3)
lead("GET /dashboard (role-aware).", "Owner payload: { role: “shop_owner”, user, shop, "
     "total_outstanding_cents, customers_count, overdue_count, repayments_this_month_cents, "
     "recent_transactions[] }. Customer payload: { role: “customer”, user, total_owed_cents, "
     "spent_this_month_cents, accounts[] (each with shop), recent_transactions[] (each with shop_name) }.")
h("8.2.4 Accounts", 3)
lead("GET /accounts?q=&filter= — owner only.", "filter ∈ {outstanding, overdue, settled} "
     "(default all). Returns { accounts: [account_summary] }.")
lead("GET /accounts/:id — owner of the shop or the linked customer.", "Returns the account detail: "
     "summary fields plus shop, invite_token, and ledger[] (each transaction with running_balance_cents). "
     "403 if unauthorized; 404 if not found.")
lead("POST /accounts — owner only.", "Request: { account: { customer_name, customer_phone } }. "
     "Success 201: account detail. Failure 422: { error, details }.")
table(["Field (account JSON)", "Type", "Meaning"],
      [["id, shop_id, customer_id", "int / null", "Identifiers (customer_id null until joined)"],
       ["customer_name, display_name", "string", "Owner-entered name; display name (user name if joined)"],
       ["customer_phone", "string/null", "Phone"],
       ["balance_cents", "int", "Cached owed amount (chetrum)"],
       ["overdue, joined", "bool", "Derived flags"],
       ["last_activity_at", "datetime/null", "Last active transaction time"],
       ["invite_token", "string", "Present in detail payloads"],
       ["ledger[]", "array", "Transactions with running_balance_cents (detail only)"]])
h("8.2.5 Transactions", 3)
lead("POST /accounts/:account_id/credits — owner only.", "Request: { transaction: { amount, "
     "description, occurred_at, itemized, notes, line_items: [{ name, quantity, unit_price }] } }. "
     "For itemized credits the total is derived from line items. Success 201: transaction JSON. Failure 422.")
lead("POST /accounts/:account_id/payments — owner only.", "Request: { transaction: { amount, "
     "payment_method, occurred_at, notes } }. Success 201: transaction JSON. Failure 422.")
lead("POST /transactions/:id/void — owner of the transaction's shop.", "Success 200: the reloaded "
     "transaction (voided=true). If already voided, 422 { error: “Transaction was already voided.” }. "
     "403 if not the owner.")
table(["Field (transaction JSON)", "Type", "Meaning"],
      [["id, account_id", "int", "Identifiers"],
       ["kind", "int", "0 credit, 1 payment"],
       ["amount_cents", "int", "Amount (chetrum)"],
       ["description, notes, payment_method", "string/null", "Optional fields"],
       ["itemized, voided", "bool", "Flags"],
       ["occurred_at", "datetime", "Business time"],
       ["created_by", "string/null", "Recorder display name"],
       ["running_balance_cents", "int/null", "Present in ledger context"],
       ["line_items[]", "array", "{ id, name, quantity, unit_price_cents, total_cents }"]])
h("8.2.6 Status codes", 3)
table(["Code", "Meaning in this API"],
      [["200 OK", "Successful read or void."],
       ["201 Created", "Resource created (session, registration, account, transaction)."],
       ["204 No Content", "Successful logout."],
       ["401 Unauthorized", "Missing or invalid bearer token; bad login credentials."],
       ["402 Payment Required", "Planned (§17): the shop's subscription is disabled; write actions are blocked until payment is approved."],
       ["403 Forbidden", "Authenticated but not permitted (e.g. non-owner, other shop, non-admin on admin endpoint)."],
       ["404 Not Found", "Record does not exist."],
       ["422 Unprocessable Entity", "Validation failure; body includes details[]."]])
h("8.3 Hardware Interfaces", 2)
para("None beyond standard client devices (smartphones, desktops/laptops) and the server host. No "
     "specialised peripherals are required in the MVP.")
h("8.4 Software Interfaces", 2)
bullet("PostgreSQL as the system of record.")
bullet("JSON API under /api/v1 consumed by the Flutter client.")
bullet("Action Mailer (SMTP/relay) for password-reset email.")
bullet("Rails health endpoint at /up for uptime checks.")
h("8.5 Communication Interfaces", 2)
bullet("HTTP/HTTPS for web and API traffic.")
bullet("Web auth via signed, httponly, same-site cookies; API auth via opaque bearer tokens.")
bullet("In development the mobile emulator reaches the backend at 10.0.2.2:3007 by default; the API base URL is configurable at build time via --dart-define=API_BASE_URL.")
doc.add_page_break()

# =====================================================================
# 9. NON-FUNCTIONAL
# =====================================================================
h("9. Non-Functional Requirements", 1)
table(["ID", "Category", "Requirement"],
      [["NFR-1", "Security – credentials", "Passwords shall be stored only as bcrypt digests (has_secure_password); plaintext passwords are never persisted or logged."],
       ["NFR-2", "Security – sessions", "Web sessions shall use signed, httponly, same-site cookies; API sessions shall use opaque per-session bearer tokens with a unique index."],
       ["NFR-3", "Security – authorization", "Account and transaction access shall be limited to the owning shop's owner or the linked customer; cross-shop access shall be impossible."],
       ["NFR-4", "Security – abuse", "Login and password-reset shall be rate-limited (10 requests / 3 minutes). Password reset shall not reveal whether an email exists."],
       ["NFR-5", "Security – invalidation", "A password reset shall destroy all of the user's existing sessions."],
       ["NFR-6", "Data integrity", "Money shall be stored only as integers; balances shall be recomputed from source transactions (never incrementally mutated). Foreign keys shall enforce referential integrity."],
       ["NFR-7", "Auditability", "All account creations, transaction creations/edits/voids, and customer joins shall produce append-only audit events with actor and metadata."],
       ["NFR-8", "Performance", "Common queries shall be index-backed (accounts by shop/customer/phone; transactions by account+occurred_at and account+kind; unique session token and invite token). Dashboards load the most recent 8 items."],
       ["NFR-9", "Reliability", "Void-and-recompute semantics shall ensure corrections never lose history; balances are recomputed within transactional callbacks."],
       ["NFR-10", "Availability", "A health endpoint (/up) shall support external monitoring; the app shall be deployable as a container."],
       ["NFR-11", "Usability", "The interface shall mirror the familiar paper khata, present amounts in Ngultrum, and show the same ledger to both parties in real time."],
       ["NFR-12", "Maintainability", "Cross-cutting logic (money, audit, authentication, serialization) shall live in shared concerns/modules used by both surfaces."],
       ["NFR-13", "Portability", "Two clients shall share one backend, data model, and domain rules; deployment shall be containerised."],
       ["NFR-14", "Compatibility", "The web app shall target modern browsers; the mobile app shall target Dart SDK 3.4+ with Material 3."],
       ["NFR-15", "Localisation-readiness", "Currency formatting shall be centralised; Dzongkha localisation is deferred but the design shall not preclude it."],
       ["NFR-16", "Privacy", "Only data necessary for the ledger shall be collected; phone numbers are optional and unverified."]])
doc.add_page_break()

# =====================================================================
# 10. MOBILE
# =====================================================================
h("10. Mobile Application Requirements (Flutter)", 1)
h("10.1 Overview", 2)
para("bangkee-app is a role-aware Flutter client that consumes the JSON API. State management uses "
     "the provider package (ChangeNotifier); there is no code generation. The theme derives a "
     "Material 3 ColorScheme from the shared design tokens.")
h("10.2 Structure", 2)
table(["Area", "Responsibility"],
      [["config.dart", "API base URL (compile-time --dart-define)"],
       ["theme/app_theme.dart", "Material 3 ColorScheme from Trust Blue tokens"],
       ["models/", "User, Account, Transaction, LineItem, Dashboard"],
       ["services/api_client.dart", "HTTP + bearer token + error mapping"],
       ["services/auth_controller.dart", "Session/token persistence (ChangeNotifier)"],
       ["services/bangkee_api.dart", "Typed endpoint wrappers"],
       ["widgets/", "BalanceCard, StatCard, TransactionTile, StatusBadge"],
       ["screens/", "auth_gate, login, register, home (role-aware nav), owner & customer dashboards, accounts, account detail, record entry, settings"]])
h("10.3 Functional expectations", 2)
fr("FR-M.1", "Token persistence",
   "The app shall persist the bearer token (shared_preferences) and route to login or home on boot via the auth gate.")
fr("FR-M.2", "Role-aware navigation",
   "The app shall present owner or customer navigation and screens according to the authenticated user's role.")
fr("FR-M.3", "Record entry",
   "Owners shall record simple/itemized credits and payments, and void transactions, from the app, mirroring web behaviour and validation.")
fr("FR-M.4", "Error mapping",
   "The API client shall map non-2xx responses to user-facing errors (401 → re-authenticate; 422 → show details).")
fr("FR-M.5", "Currency formatting",
   "The app shall format chetrum amounts as Ngultrum locally (util/money.dart, intl).")
doc.add_page_break()

# =====================================================================
# 11. VALIDATION
# =====================================================================
h("11. Validation & Error-Handling Catalogue", 1)
table(["Field / rule", "Validation", "On failure"],
      [["User.email_address", "Present, unique, normalized lower-case", "422 / inline form error"],
       ["User.name", "Present", "422 / inline form error"],
       ["User.role", "In {shop_owner, customer}", "422 / inline form error"],
       ["User.password", "Confirmed (web) via password_confirmation", "Re-render form / re-prompt"],
       ["Shop.name", "Present", "422 / inline form error"],
       ["Shop.credit_due_days", "Numeric, > 0", "422 / inline form error"],
       ["Account.customer_name", "Present", "422 / inline form error"],
       ["Account.customer_phone", "Normalized (blank → null)", "Stored as null"],
       ["Transaction.amount_cents", "Numeric, > 0", "422 / inline form error"],
       ["Transaction.occurred_at", "Present (defaults to now on create)", "422"],
       ["Transaction.kind", "In {credit, payment}", "422"],
       ["LineItem.name", "Present", "422 / inline error"],
       ["LineItem.quantity", "Numeric, > 0", "422 / inline error"],
       ["LineItem.unit_price_cents", "Numeric, >= 0", "422 / inline error"],
       ["Itemized credit", "At least one valid line item; total derived", "Blank rows rejected"],
       ["Void on voided txn", "Must not already be voided", "No-op; alert / 422"],
       ["Invalid bearer token", "Token must match a session", "401 Unauthorized"],
       ["Non-owner on owner action", "Role must be shop_owner", "403 / redirect with alert"],
       ["Cross-shop / foreign account", "Owner must own the shop, or be the linked customer", "403 / redirect with alert"],
       ["Reset token", "Valid signature, unexpired", "Redirect with alert"]])
doc.add_page_break()

# =====================================================================
# 12. SECURITY
# =====================================================================
h("12. Security Requirements", 1)
h("12.1 Authentication", 2)
bullet("Passwords are hashed with bcrypt via has_secure_password; authentication uses User.authenticate_by to mitigate timing attacks.")
bullet("Web sessions are tracked by a signed, httponly, same-site permanent cookie referencing a Session row.")
bullet("API sessions are identified by an opaque token (unique-indexed) presented as a bearer token.")
bullet("Login and password-reset endpoints are rate-limited to 10 requests per 3 minutes.")
h("12.2 Authorization", 2)
para("Authorization is enforced server-side on every request. The matrix in Appendix B defines "
     "permitted actions by role. Account-level access is restricted to the shop owner or the linked "
     "customer; transaction edit/void is restricted to the owner of the transaction's shop.")
h("12.3 Session & account safety", 2)
bullet("A password reset destroys all existing sessions for the user.")
bullet("Password-reset requests do not disclose whether an email is registered.")
bullet("Reset tokens are signed and expiring; tampered or expired tokens are rejected.")
bullet("Sign-out destroys the current session and clears the cookie.")
h("12.4 Transport & platform", 2)
bullet("All traffic is expected over HTTPS in production.")
bullet("The web app restricts itself to modern browsers, reducing the supported attack surface.")
bullet("Input is strong-parameter filtered (params.expect) on every controller action.")
doc.add_page_break()

# =====================================================================
# 13. DATA & SEED
# =====================================================================
h("13. Data Management, Migration & Seed Data", 1)
h("13.1 Migrations", 2)
para("The schema is created by ordered migrations: users, sessions, shops, accounts, transactions, "
     "line_items, audit_events, and a later migration adding the session token. The authoritative "
     "schema is db/schema.rb (version 2026_06_22_070000).")
h("13.2 Seed data", 2)
para("db/seeds.rb is idempotent and provisions a demo shop and customers for development and demonstration:")
bullet("Owner Karma Wangmo (karma@shop.bt) owning “Karma General Shop” (credit_due_days 30).")
bullet("Customer Dawa Tshering (dawa@example.bt), linked to an account, plus unlinked accounts Sonam Dorji and Pema Lhamo.")
bullet("Sample transactions: a simple credit, an itemized credit, a payment, an overdue account (50-day-old credit), and a settled account.")
bullet("Demo logins use the password “password”.")
h("13.3 Retention", 2)
bullet("Transactions are voided rather than deleted in normal operation; voided rows are retained for audit.")
bullet("Audit events are append-only.")
bullet("Destroying an account cascades to its transactions, their line items, and related audit events.")
doc.add_page_break()

# =====================================================================
# 14. TESTING
# =====================================================================
h("14. Testing & Acceptance Criteria", 1)
para("The system is verified by model, controller, and integration tests (Rails Minitest). The "
     "following acceptance criteria, derived from the existing test suite and requirements, define "
     "“done” for the core ledger.")
h("14.1 Ledger & balance", 2)
bullet("A credit increases the cached balance by its amount; a payment decreases it (verified: 1,250 → 125,000 chetrum; credit 1,000 then payment 400 → 60,000).")
bullet("An itemized credit's amount equals the sum of its line-item totals (2×600 + 1×450 → 165,000 chetrum).")
bullet("The ledger reports running balances in chronological order (credits 1,000 & 500 then payment 600 → 100,000 / 150,000 / 90,000).")
bullet("signed_amount_cents is positive for credits and negative for payments.")
h("14.2 Void & audit", 2)
bullet("Voiding removes the amount from the balance and records exactly one audit event.")
bullet("Voided transactions are excluded from the ledger and balance.")
bullet("Voiding an already-voided transaction is a no-op (returns false).")
h("14.3 Overdue", 2)
bullet("An account with a positive balance and no recent activity (60-day-old credit) is overdue.")
bullet("A settled account (credit then equal payment) is not overdue.")
h("14.4 Validation & search", 2)
bullet("A transaction with a non-positive amount is invalid (error on amount_cents).")
bullet("Account search matches both customer name and phone and returns empty for non-matches.")
h("14.5 Flows", 2)
bullet("Dashboard integration: role-aware dashboards render the expected aggregates.")
bullet("API integration: the documented endpoints authenticate, authorize, and return the specified payloads and status codes.")
bullet("Controller tests cover accounts, credits, payments, transactions, sessions, and passwords.")
doc.add_page_break()

# =====================================================================
# 15. CONSTRAINTS/RISKS
# =====================================================================
h("15. Constraints, Assumptions & Risks", 1)
h("15.1 Key constraints (recap)", 2)
bullet("Single shop per owner; multi-shop deferred.")
bullet("Online-first; no offline operation in the MVP.")
bullet("Currency limited to Ngultrum; Dzongkha localisation deferred.")
bullet("The LineItem→:purchase association name must be preserved.")
h("15.2 Risks & mitigations", 2)
table(["Risk", "Mitigation"],
      [["Balance drift from concurrent writes", "Always recompute from source transactions in model callbacks rather than incrementing."],
       ["Unauthorized data access", "Server-side authorization on every action; shop-scoping; bearer-token checks."],
       ["Lost history from corrections", "Void instead of delete; append-only audit trail."],
       ["Connectivity gaps", "Documented as an assumption; offline sync is a planned enhancement."],
       ["Account enumeration via reset", "Reset responses do not reveal account existence."]])
doc.add_page_break()

# =====================================================================
# 16. CUSTOMER PAYMENTS VIA BANK TRANSFER (PLANNED)
# =====================================================================
h("16. Customer Payments via Bank Transfer (Proof & Confirmation) — Planned", 1)
lead("Status: Planned.", "This module is specified as a committed requirement for the next "
     "release; it is not part of the as-built 2.0 baseline. It depends on file attachments "
     "(Active Storage — image_processing is already in the Gemfile, but the Active Storage tables "
     "are not yet migrated) and on the in-app notification capability introduced in §16.7.")

h("16.1 Overview", 2)
para("Today only a shop owner can record a payment. This module lets a customer settle their "
     "balance by bank transfer (or mobile wallet) and prove it: the owner publishes the shop's bank "
     "details, the customer transfers the money outside Bangkee and uploads a screenshot as proof, "
     "the owner is notified, reviews the screenshot, and confirms — at which point Bangkee records "
     "a real payment transaction that reduces the balance. A proof never changes the balance until "
     "the owner confirms it, preserving the existing single-source-of-truth ledger.")
h("16.2 Actors & Trigger", 2)
bullet("Shop owner — maintains bank details; reviews and confirms/rejects proofs.")
bullet("Customer (linked to the account) — views bank details, submits a payment proof with a screenshot.")
bullet("Trigger — a customer who has paid by transfer wants the shop to acknowledge and settle it.")

h("16.3 Proposed Data Model", 2)
para("Two new entities are introduced. Amounts remain integer chetrum; screenshots are Active "
     "Storage attachments.")
para("bank_accounts (a shop's published payment destinations)", bold=True)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["shop_id", "bigint", "no", "—", "FK → shops"],
       ["bank_name", "string", "no", "—", "e.g. Bank of Bhutan, BNB"],
       ["account_name", "string", "no", "—", "Account holder name; presence validated"],
       ["account_number", "string", "no", "—", "Bank account number; presence validated"],
       ["branch", "string", "yes", "—", "Optional branch"],
       ["mobile_wallet", "string", "yes", "—", "Optional mobile-wallet number (e.g. mPay/MyPay)"],
       ["qr_image", "attachment", "yes", "—", "Optional payment-QR image (Active Storage)"],
       ["instructions", "text", "yes", "—", "Free-text payment instructions"],
       ["primary", "boolean", "no", "false", "Marks the default destination shown to customers"],
       ["active", "boolean", "no", "true", "Hidden from customers when false"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("payment_proofs (customer-submitted settlement claims)", bold=True)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["account_id", "bigint", "no", "—", "FK → accounts"],
       ["submitted_by_id", "bigint", "no", "—", "FK → users; the customer"],
       ["amount_cents", "bigint", "no", "0", "Claimed amount (chetrum); must be > 0"],
       ["reference", "string", "yes", "—", "Bank/wallet reference or journal no."],
       ["note", "text", "yes", "—", "Optional customer note"],
       ["screenshot", "attachment", "no", "—", "Proof image (Active Storage); required"],
       ["status", "integer", "no", "0", "Enum: 0 pending, 1 confirmed, 2 rejected"],
       ["reviewed_by_id", "bigint", "yes", "—", "FK → users; the owner who reviewed"],
       ["reviewed_at", "datetime", "yes", "—", "When confirmed/rejected"],
       ["rejection_reason", "string", "yes", "—", "Required when rejected"],
       ["transaction_id", "bigint", "yes", "—", "FK → transactions; the payment created on confirm"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])

h("16.4 Business Rules", 2)
table(["ID", "Rule"],
      [["BR-27", "An owner may publish zero or more bank accounts for their shop; at most one is primary. account_name and account_number are required."],
       ["BR-28", "A customer may submit a payment proof only for an account they are linked to (customer_id = the user). amount_cents must be > 0 and a screenshot is required."],
       ["BR-29", "A new proof is created with status = pending and MUST NOT change the account balance."],
       ["BR-30", "On confirm by the shop owner, the system creates an active Payment transaction (kind = payment) for the proof's amount, sets the proof to confirmed, links transaction_id, and recomputes the balance. The action is atomic."],
       ["BR-31", "On reject by the shop owner, the proof is set to rejected with a reason; no transaction is created and the balance is unchanged."],
       ["BR-32", "confirmed and rejected are terminal states; a reviewed proof cannot be re-reviewed or edited. Confirming/rejecting an already-reviewed proof is a no-op."],
       ["BR-33", "Only the owner of the proof's shop may review it; the screenshot is viewable only by that owner and by the submitting customer."],
       ["BR-34", "Confirming a proof for more than the outstanding balance is permitted and may drive the balance negative (in credit), consistent with BR-5."],
       ["BR-35", "A new pending proof notifies the shop owner (§16.7); confirmation/rejection notifies the customer."]])

h("16.5 Functional Requirements", 2)
fr("FR-11.1", "Manage shop bank details",
   "An owner shall add, edit, deactivate, and delete bank accounts (and an optional payment-QR image and instructions) for their shop, and mark one as primary.",
   validation="account_name and account_number present; at most one primary; image type/size limits on qr_image.")
fr("FR-11.2", "View shop bank details (customer)",
   "A linked customer shall view the shop's active bank accounts and instructions so they can transfer funds.")
fr("FR-11.3", "Submit a payment proof",
   "A linked customer shall submit a payment proof with an amount, an optional reference/note, and a required screenshot.",
   inputs="amount, reference, note, screenshot (image).",
   processing="Create a payment_proof with status pending; attach the screenshot; notify the shop owner.",
   validation="amount_cents > 0; screenshot present and an accepted image type within the size limit; submitter linked to the account.",
   errors="Non-linked user → 403; missing/oversized/invalid screenshot → 422.")
fr("FR-11.4", "Review pending proofs",
   "An owner shall see a list of pending proofs for their shop with the screenshot, amount, reference, customer, and submission time.")
fr("FR-11.5", "Confirm a payment proof",
   "An owner shall confirm a pending proof, which atomically records a payment transaction for the amount, links it, marks the proof confirmed, recomputes the balance, and notifies the customer.",
   errors="Reviewing an already-reviewed proof is a no-op (alert / 422).")
fr("FR-11.6", "Reject a payment proof",
   "An owner shall reject a pending proof with a reason; no transaction is created and the customer is notified.",
   validation="rejection_reason required on reject.")
fr("FR-11.7", "Proof history",
   "Both parties shall see the proof's lifecycle (pending → confirmed/rejected) on the account; a confirmed proof links to its payment transaction.",
   priority="Medium")

h("16.6 Use Cases", 2)
use_case("UC-13", "Owner publishes bank details", "Shop owner",
         "Owner wants customers to pay by transfer.", "Authenticated owner.", "Owner opens shop payment settings.",
         ["Owner enters bank name, account name and number, optional branch/wallet/QR/instructions.",
          "Owner marks one account primary and saves.",
          "System validates and stores the bank account; it becomes visible to linked customers."],
         exceptions=["1a. Account name/number blank → validation error."],
         post=["The shop has published, customer-visible bank details."], freq="Rare")
use_case("UC-14", "Customer submits a payment proof", "Customer",
         "Customer paid by transfer and wants it settled.",
         "Authenticated customer linked to the account; owner has published bank details.",
         "Customer chooses “I've paid — upload proof.”",
         ["Customer views the shop's bank details and transfers funds externally.",
          "Customer enters the amount and optional reference/note and uploads a screenshot.",
          "System creates a pending proof and notifies the owner.",
          "System confirms submission and shows it as pending in history."],
         exceptions=["2a. No screenshot or invalid image → 422.", "2b. Not linked to the account → 403."],
         post=["A pending proof exists; balance unchanged until the owner confirms."], freq="Frequent")
use_case("UC-15", "Owner confirms or rejects a proof", "Shop owner",
         "Owner verifies the transfer actually arrived.", "A pending proof exists for the owner's shop.",
         "Owner opens the pending proofs list.",
         ["Owner views the screenshot, amount, and reference.",
          "If the transfer is genuine, owner confirms → a payment transaction is recorded, the balance drops, the proof is confirmed, and the customer is notified.",
          "If not, owner rejects with a reason → no transaction; customer is notified."],
         alts=["2a. Amount on the screenshot differs → owner may reject and ask the customer to resubmit."],
         exceptions=["1a. Proof already reviewed → no-op."],
         post=["Confirmed: balance reduced and proof linked to its payment. Rejected: balance unchanged."], freq="Frequent")

h("16.7 Notifications", 2)
bullet("In-app notification to the owner on each new pending proof (count badge on the pending-proofs surface).")
bullet("In-app notification to the customer on confirm/reject, with the outcome and (on reject) the reason.")
bullet("Email notification SHOULD accompany in-app events, reusing Action Mailer; push/SMS remain deferred (§18).")

h("16.8 API Additions", 2)
table(["Method & path", "Purpose", "Auth"],
      [["GET /api/v1/bank_accounts", "List the shop's bank accounts", "Owner"],
       ["POST /api/v1/bank_accounts", "Add a bank account", "Owner"],
       ["PATCH /api/v1/bank_accounts/:id", "Edit a bank account", "Owner"],
       ["DELETE /api/v1/bank_accounts/:id", "Remove a bank account", "Owner"],
       ["GET /api/v1/accounts/:id (extended)", "Account detail now includes the shop's active bank accounts", "Owner/linked customer"],
       ["POST /api/v1/accounts/:account_id/payment_proofs", "Submit a proof (multipart: amount, reference, note, screenshot)", "Linked customer"],
       ["GET /api/v1/accounts/:account_id/payment_proofs", "List proofs for the account", "Owner/linked customer"],
       ["POST /api/v1/payment_proofs/:id/confirm", "Confirm → records a payment", "Owner"],
       ["POST /api/v1/payment_proofs/:id/reject", "Reject with a reason", "Owner"]])

h("16.9 Edge Cases & Error Handling", 2)
bullet("A screenshot exceeding the size limit or of a disallowed type MUST be rejected at upload (422) with no proof created.")
bullet("Duplicate proofs for the same transfer MUST be reviewable independently; the owner rejects duplicates.")
bullet("Concurrent confirms of the same proof MUST be idempotent — exactly one payment transaction is created.")
bullet("Voiding the payment transaction created from a proof MUST follow normal void rules and SHOULD mark the linked proof as voided/needs-review in history.")
bullet("If notification delivery fails, the proof MUST still be persisted and visible in the pending list (delivery is retried; never a silent drop).")
bullet("A confirmed proof's screenshot MUST be retained for audit even after the payment is voided.")

doc.add_page_break()

# =====================================================================
# 17. SUBSCRIPTION, BILLING & PLAN ENFORCEMENT (PLANNED)
# =====================================================================
h("17. Subscription, Billing & Plan Enforcement — Planned", 1)
lead("Status: Planned.", "This module is specified as a committed requirement for the next "
     "release; it is not part of the as-built 2.0 baseline. It introduces the Platform "
     "Administrator user class (§2.3) and depends on file attachments (Active Storage) and the "
     "in-app/email notification capability (§16.7).")

h("17.1 Overview", 2)
para("Bangkee is sold to shop owners as a subscription at Nu 200 per month. In this first, "
     "manual-billing version there is no payment gateway: the shop owner pays the Bangkee operator "
     "(the platform administrator) directly by bank transfer or mobile wallet and submits a payment "
     "request with a screenshot; the platform administrator reviews it and approves or rejects. An "
     "approval extends the shop's paid period by one month. When a shop's paid period lapses, the "
     "owner is shown escalating warnings; after a two-week grace period without an approved payment, "
     "the shop's write features are disabled until payment is approved. Customers are never billed "
     "and their read access to their own ledgers is never disabled.")
h("17.2 Plan & Pricing", 2)
bullet("Single plan: “Standard”, Nu 200 / month (price_cents = 20,000; configurable).")
bullet("New shops begin on a free trial (trial_days, default 14) so owners can evaluate before paying.")
bullet("Billing is per shop (one shop per owner in this version).")

h("17.3 Proposed Data Model", 2)
para("A platform_admin boolean is added to users to designate the operator. Two new entities track "
     "the subscription and the manual payment requests.")
para("users (addition)", bold=True)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["platform_admin", "boolean", "no", "false", "True only for the Bangkee operator account(s)"]])
para("subscriptions (one per shop)", bold=True)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["shop_id", "bigint", "no", "—", "FK → shops; unique (one per shop)"],
       ["plan", "string", "no", "standard", "Plan code"],
       ["price_cents", "bigint", "no", "20000", "Monthly price (Nu 200)"],
       ["status", "integer", "no", "0", "Enum: 0 trialing, 1 active, 2 past_due, 3 disabled"],
       ["current_period_end", "datetime", "no", "—", "Paid-through date (trial end initially)"],
       ["grace_until", "datetime", "yes", "—", "current_period_end + 14 days; end of usable grace"],
       ["disabled_at", "datetime", "yes", "—", "When write features were locked"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])
para("subscription_payments (manual payment requests to the operator)", bold=True)
table(["Column", "Type", "Null", "Default", "Notes"],
      [["id", "bigint", "no", "—", "Primary key"],
       ["shop_id", "bigint", "no", "—", "FK → shops"],
       ["submitted_by_id", "bigint", "no", "—", "FK → users; the shop owner"],
       ["amount_cents", "bigint", "no", "0", "Amount paid (chetrum); must be > 0"],
       ["method", "string", "yes", "—", "Bank transfer / mobile wallet"],
       ["reference", "string", "yes", "—", "Transfer reference / journal no."],
       ["screenshot", "attachment", "no", "—", "Proof image (Active Storage); required"],
       ["months", "integer", "no", "1", "Months of cover requested; must be >= 1"],
       ["status", "integer", "no", "0", "Enum: 0 pending, 1 approved, 2 rejected"],
       ["reviewed_by_id", "bigint", "yes", "—", "FK → users; the platform admin"],
       ["reviewed_at", "datetime", "yes", "—", "When approved/rejected"],
       ["rejection_reason", "string", "yes", "—", "Required when rejected"],
       ["created_at / updated_at", "datetime", "no", "—", "Timestamps"]])

h("17.4 Business Rules", 2)
table(["ID", "Rule"],
      [["BR-36", "Every shop has exactly one subscription, created on shop provisioning with status trialing and current_period_end = now + trial_days."],
       ["BR-37", "Subscription status derives from server time: active/trialing while now <= current_period_end; past_due while current_period_end < now <= grace_until; disabled once now > grace_until without an approved payment."],
       ["BR-38", "grace_until = current_period_end + 14 days. During trialing, active, and past_due, all owner features remain fully usable."],
       ["BR-39", "In past_due, the owner UI MUST show an escalating warning banner (days remaining in grace, amount due) on every screen."],
       ["BR-40", "In disabled, owner WRITE actions MUST be blocked — add/edit/remove customer, record credit/payment, void, edit transaction, manage bank details, and confirm customer payment proofs (§16). The owner retains read access and access to the billing screen to pay."],
       ["BR-41", "Customers are never billed and never see subscription UI; customer read access to their own ledgers is NEVER disabled by a shop's billing state."],
       ["BR-42", "Only a platform_admin may approve or reject subscription payments, view all shops' billing status, or manually change a subscription state."],
       ["BR-43", "Approving a subscription_payment extends current_period_end by (months × 1 month) from the later of now or the existing current_period_end, recomputes grace_until, clears disabled_at, sets status to active, and notifies the owner. The action is atomic."],
       ["BR-44", "Rejecting a subscription_payment changes nothing about the period; the owner is notified with the reason and may resubmit."],
       ["BR-45", "approved and rejected are terminal states for a subscription_payment; re-review is a no-op."],
       ["BR-46", "All subscription state changes and review actions are recorded as audit events (§ Appendix D)."]])

h("17.5 Functional Requirements", 2)
fr("FR-12.1", "Subscription provisioning",
   "On shop creation the system shall create a subscription in trialing status with current_period_end = now + trial_days.")
fr("FR-12.2", "View subscription status",
   "An owner shall view their shop's subscription status, paid-through date, amount due, and the operator's payment instructions (bank/wallet details).")
fr("FR-12.3", "Submit a subscription payment",
   "An owner shall submit a payment request with an amount, optional method/reference, months of cover, and a required screenshot.",
   inputs="amount, method, reference, months, screenshot.",
   processing="Create a subscription_payment with status pending; notify platform administrators.",
   validation="amount_cents > 0; months >= 1; screenshot present, accepted type within size limit.",
   errors="Missing/invalid screenshot → 422.")
fr("FR-12.4", "Warning on lapse",
   "When a subscription is past_due, the system shall display an escalating warning banner across the owner UI showing days remaining before lockout and the amount due.")
fr("FR-12.5", "Feature lockout",
   "When a subscription is disabled (grace exhausted), the system shall block all owner write actions and return HTTP 402 on write API endpoints, while preserving read access and access to the billing screen.",
   errors="Blocked web actions redirect to the billing screen with an explanation; blocked API writes return 402.")
fr("FR-12.6", "Admin review queue",
   "A platform administrator shall view pending subscription payments across all shops with the screenshot, amount, months, shop, and submitter.")
fr("FR-12.7", "Approve a subscription payment",
   "A platform administrator shall approve a pending payment, atomically extending the paid period, clearing any disabled state, and notifying the owner.",
   errors="Reviewing an already-reviewed payment is a no-op (422).")
fr("FR-12.8", "Reject a subscription payment",
   "A platform administrator shall reject a pending payment with a reason; the period is unchanged and the owner is notified.",
   validation="rejection_reason required on reject.")
fr("FR-12.9", "Admin shop overview",
   "A platform administrator shall view all shops with their subscription status, paid-through date, and last payment, and may manually adjust a subscription state for support purposes (audited).",
   priority="Medium")
fr("FR-12.10", "Reactivation",
   "Approval of a payment for a disabled shop shall immediately restore full owner functionality.")

h("17.6 Use Cases", 2)
use_case("UC-16", "Owner pays the monthly subscription", "Shop owner",
         "Owner wants to keep using Bangkee.", "Authenticated owner; subscription past_due or nearing expiry.",
         "Owner opens the billing screen.",
         ["Owner reads the operator's payment instructions and transfers Nu 200 externally.",
          "Owner enters the amount, months, optional reference, and uploads a screenshot.",
          "System creates a pending subscription payment and notifies the platform administrator.",
          "System shows the request as pending."],
         exceptions=["2a. No screenshot/invalid → 422."],
         post=["A pending subscription payment awaits admin review; status unchanged until approved."], freq="Monthly")
use_case("UC-17", "Administrator approves or rejects a payment", "Platform administrator",
         "Operator must record revenue and extend access correctly.",
         "A pending subscription payment exists.", "Admin opens the review queue.",
         ["Admin views the screenshot, amount, months, and shop.",
          "If genuine, admin approves → the paid period extends by the months covered, any lockout is cleared, status becomes active, and the owner is notified.",
          "If not, admin rejects with a reason → period unchanged; owner notified."],
         exceptions=["1a. Payment already reviewed → no-op."],
         post=["Approved: shop paid-through advanced and re-enabled. Rejected: unchanged."], freq="Daily (operator)")
use_case("UC-18", "Lapse → warning → lockout", "System / Shop owner",
         "Operator wants to encourage payment without abruptly cutting access.",
         "A subscription's current_period_end has passed.", "Server time crosses the period end.",
         ["On expiry, the subscription becomes past_due and the owner sees a warning banner with days left in grace.",
          "The owner is reminded (in-app/email) during the 14-day grace.",
          "If no payment is approved by grace_until, the subscription becomes disabled and owner write actions are blocked.",
          "The owner can still view data and open the billing screen; on an approved payment, access is restored immediately."],
         alts=["2a. Owner pays during grace → approval returns the subscription to active with no lockout."],
         post=["Access reflects payment state; customers are unaffected throughout."], freq="As needed")

h("17.7 Notifications", 2)
bullet("Platform administrator notified (in-app/email) on each new pending subscription payment.")
bullet("Owner notified on approval (new paid-through date) and on rejection (with reason).")
bullet("Owner reminded on entering past_due and at intervals during the grace period; a final notice on lockout.")

h("17.8 API Additions", 2)
table(["Method & path", "Purpose", "Auth"],
      [["GET /api/v1/subscription", "Own shop's subscription status & instructions", "Owner"],
       ["POST /api/v1/subscription/payments", "Submit a payment request (multipart w/ screenshot)", "Owner"],
       ["GET /api/v1/subscription/payments", "List own shop's payment requests", "Owner"],
       ["GET /api/v1/admin/subscription_payments?status=pending", "Review queue", "Platform admin"],
       ["POST /api/v1/admin/subscription_payments/:id/approve", "Approve & extend", "Platform admin"],
       ["POST /api/v1/admin/subscription_payments/:id/reject", "Reject with reason", "Platform admin"],
       ["GET /api/v1/admin/shops", "All shops with billing status", "Platform admin"]])
para("Enforcement: when a shop is disabled, all owner write endpoints (account/credit/payment/void/"
     "bank-account/proof-confirm) return 402 Payment Required; read endpoints and the billing "
     "endpoints remain available.")

h("17.9 Edge Cases & Error Handling", 2)
bullet("All state derivation uses authoritative server time; client clocks never affect status (BR-37).")
bullet("A disabled owner MUST still reach the billing screen and read-only views; only writes are blocked.")
bullet("Approving overlapping or duplicate payments extends from current_period_end, not from now, so cover is never lost or double-applied; the admin rejects clear duplicates.")
bullet("At the exact grace boundary, status MUST resolve deterministically (now > grace_until ⇒ disabled).")
bullet("A platform_admin account MUST NOT also act as a shop owner or customer on the same identity; admin endpoints reject non-admins with 403.")
bullet("Subscription screenshots MUST be validated for type/size like §16 proofs.")
bullet("If a shop is deleted, its subscription and payment history are removed; revenue records SHOULD be retained per the operator's accounting policy.")
bullet("Notification failures MUST NOT block state changes; reminders are retried.")

doc.add_page_break()

# =====================================================================
# 18. FUTURE
# =====================================================================
h("18. Future Enhancements (Out of Scope for the MVP)", 1)
para("The following capabilities are explicitly deferred beyond the current baseline:")
bullet("Push / SMS notifications (e.g. overdue and payment alerts) — a push-notification mockup exists in design/.")
bullet("Real-time streaming via Action Cable.")
bullet("Turbo Native shells for mobile.")
bullet("Offline operation and synchronisation.")
bullet("CSV / statement export and printing.")
bullet("Dzongkha localisation (i18n).")
bullet("Multiple shops per owner and staff/cashier sub-roles.")
doc.add_page_break()

# =====================================================================
# APPENDICES
# =====================================================================
h("Appendix A — Requirements Traceability Matrix", 1)
table(["Requirement group", "Primary implementing components", "Verified by"],
      [["FR-1 Auth & registration", "Authentication concern; Sessions/Registrations/Passwords controllers; User, Session models", "sessions/passwords controller tests; user test"],
       ["FR-2 Account management", "AccountsController; Account model (search/outstanding/settled scopes)", "accounts controller test; account test"],
       ["FR-3/4 Credits & payments", "Credits/Payments controllers; Transaction & LineItem models; HasMoneyAttribute", "credits/payments controller tests; transaction test"],
       ["FR-5 Void & recompute", "Transaction#void!; Account#recompute_balance!; Auditable", "transaction test; account test"],
       ["FR-6 Dashboards", "DashboardsController; Api::V1::DashboardController", "dashboard_flow integration test"],
       ["FR-7 Statements & ledger", "StatementsController; Account#ledger", "account test (ledger)"],
       ["FR-8 Invitations", "InvitationsController; Registrations claim logic; Account invite_token", "registrations/controller coverage"],
       ["FR-9 Overdue", "Shop#overdue_accounts; Account#overdue?; credit_due_days", "account test (overdue)"],
       ["FR-10 JSON API", "Api::V1::* controllers; BaseController; ApiSerializers", "api/v1 api_flow integration test"],
       ["FR-M Mobile", "bangkee-app services/screens", "manual / app-level testing"],
       ["BR-1…BR-26 Domain rules", "Models & concerns (HasMoneyAttribute, Auditable, Transaction, Account, Shop)", "model tests"],
       ["FR-11 / BR-27…BR-35 Bank-transfer settlement (Planned, §16)", "Planned: BankAccount, PaymentProof models; Active Storage; owner review controller/screens", "Not yet implemented — acceptance per §16"],
       ["FR-12 / BR-36…BR-46 Subscription billing (Planned, §17)", "Planned: Subscription, SubscriptionPayment models; platform_admin; enforcement before_action; admin controllers/screens", "Not yet implemented — acceptance per §17"]])
doc.add_page_break()

h("Appendix B — Authorization Matrix", 1)
table(["Action", "Shop owner (own shop)", "Customer (own account)", "Other / unauth"],
      [["View own dashboard", "Yes", "Yes", "No"],
       ["List/search accounts", "Yes", "No", "No"],
       ["Create account", "Yes", "No", "No"],
       ["Edit/remove account", "Yes", "No", "No"],
       ["View account detail/ledger", "Yes", "Yes (own)", "No"],
       ["View statement", "Yes", "Yes (own)", "No"],
       ["Record credit/payment", "Yes", "No", "No"],
       ["Edit/void transaction", "Yes", "No", "No"],
       ["Join via invitation", "No (refused)", "Yes", "Visitor → sign-up"],
       ["Access another shop's data", "No", "No", "No"],
       ["Manage shop bank details (§16, Planned)", "Yes", "No", "No"],
       ["View shop bank details (§16, Planned)", "Yes", "Yes (own account)", "No"],
       ["Submit payment proof (§16, Planned)", "No", "Yes (own account)", "No"],
       ["Confirm/reject payment proof (§16, Planned)", "Yes", "No", "No"],
       ["View own subscription / pay (§17, Planned)", "Yes", "n/a (not billed)", "No"],
       ["Approve/reject subscription payments (§17, Planned)", "No (unless platform admin)", "No", "Platform admin only"],
       ["View all shops' billing (§17, Planned)", "No", "No", "Platform admin only"]])
para("Note: in the Planned subscription module (§17) a shop in the disabled state retains all read "
     "access shown above; only owner write actions are blocked. Customer read access is never "
     "affected by a shop's billing state.")
doc.add_page_break()

h("Appendix C — Glossary of Domain Terms", 1)
para("See Section 1.4 for the primary definitions. Additional clarifications:")
kv_bullet("Settled account", "balance_cents == 0 — the customer owes nothing.")
kv_bullet("In-credit account", "balance_cents < 0 — the shop effectively holds a customer prepayment/overpayment.")
kv_bullet("Active transaction", "A transaction with voided_at == null; only active transactions affect balance and ledger.")
kv_bullet("Display name", "For a joined account, the linked user's name; otherwise the owner-entered customer_name.")
doc.add_page_break()

h("Appendix D — Audit Event Catalogue", 1)
table(["action", "auditable", "When written", "Metadata"],
      [["account_created", "Account", "Owner creates a customer account", "{ customer_name }"],
       ["created", "Transaction", "Credit or payment recorded", "{ kind, amount_cents }"],
       ["edited", "Transaction", "Transaction updated (web edit)", "{ amount_cents }"],
       ["voided", "Transaction", "Transaction voided", "{ amount_cents, kind }"],
       ["customer_joined", "Account", "Customer links via invitation/registration", "{ user_id }"],
       ["bank_account_added (§16, Planned)", "Shop", "Owner publishes bank details", "{ bank_name }"],
       ["payment_proof_submitted (§16, Planned)", "PaymentProof", "Customer uploads a proof", "{ amount_cents }"],
       ["payment_proof_confirmed (§16, Planned)", "PaymentProof", "Owner confirms a proof", "{ amount_cents, transaction_id }"],
       ["payment_proof_rejected (§16, Planned)", "PaymentProof", "Owner rejects a proof", "{ amount_cents, reason }"],
       ["subscription_payment_submitted (§17, Planned)", "SubscriptionPayment", "Owner submits a payment request", "{ amount_cents, months }"],
       ["subscription_approved (§17, Planned)", "Subscription", "Admin approves; period extended", "{ amount_cents, new_period_end }"],
       ["subscription_rejected (§17, Planned)", "SubscriptionPayment", "Admin rejects a payment", "{ amount_cents, reason }"],
       ["subscription_disabled (§17, Planned)", "Subscription", "Grace exhausted; features locked", "{ grace_until }"],
       ["subscription_reactivated (§17, Planned)", "Subscription", "Payment restores access", "{ new_period_end }"]])
doc.add_page_break()

h("Appendix E — State-Transition Tables", 1)
h("E.1 Transaction lifecycle", 2)
table(["State", "Event", "Next state", "Effect"],
      [["(none)", "create", "Active", "Balance recomputed; “created” audit"],
       ["Active", "edit", "Active", "Balance recomputed; “edited” audit"],
       ["Active", "void", "Voided", "Excluded from balance/ledger; “voided” audit"],
       ["Voided", "void (again)", "Voided", "No-op (returns false)"],
       ["Active/Voided", "destroy", "(removed)", "Balance recomputed"]])
h("E.2 Account linking lifecycle", 2)
table(["State", "Event", "Next state", "Effect"],
      [["Unlinked (customer_id null)", "customer joins via token", "Linked", "customer_id set; phone backfilled; “customer_joined” audit"],
       ["Linked", "another customer opens token", "Linked", "Refused (already claimed)"],
       ["Unlinked", "shop owner opens token", "Unlinked", "Refused with guidance"]])
doc.add_page_break()

h("Appendix F — Open Issues", 1)
table(["#", "Issue", "Status"],
      [["1", "Notification channel (push/SMS) selection and provider", "Deferred (see §18)"],
       ["2", "Multi-shop ownership and staff roles", "Deferred"],
       ["3", "Statement export format (PDF/CSV)", "Deferred"],
       ["4", "Dzongkha localisation scope", "Deferred"]])

out = "/Users/nimayonten/personal/bangkee/bangkee-srs.docx"
doc.save(out)
print("Wrote", out)
